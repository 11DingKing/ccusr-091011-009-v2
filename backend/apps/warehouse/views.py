"""
仓库管理视图
"""
import logging
import io
import datetime
from decimal import Decimal
from django.http import HttpResponse
from django.utils import timezone
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from apps.core.response import success_response, error_response
from .models import (
    Unit, Category, Variety, Goods, StockIn, StockOut, Warning, Approval,
    ReviewRuleVersion, ReviewTask, ReviewRecord,
)
from .serializers import (
    UnitSerializer, UnitCreateSerializer,
    CategorySerializer, CategoryCreateSerializer,
    VarietySerializer, VarietyCreateSerializer,
    GoodsSerializer, StockInSerializer, StockOutSerializer,
    WarningSerializer, ApprovalSerializer,
    ReviewRuleSerializer, ReviewRulePublishSerializer,
    ReviewTaskSerializer, ReviewRecordSerializer, ReviewCompleteSerializer,
)
from . import review as review_service

logger = logging.getLogger('apps')


# ==================== 单位管理 ====================

class UnitListView(APIView):
    """单位列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Unit.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        units = queryset[start:end]
        
        serializer = UnitSerializer(units, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建单位"""
        serializer = UnitCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit = Unit.objects.create(
            name=serializer.validated_data['name'],
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created unit {unit.name}")
        
        return success_response(data=UnitSerializer(unit).data, message='创建成功')


class UnitDetailView(APIView):
    """单位详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新单位"""
        try:
            unit = Unit.objects.get(pk=pk)
        except Unit.DoesNotExist:
            return error_response(message='单位不存在', code=404)
        
        serializer = UnitCreateSerializer(data=request.data, context={'instance': unit})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit.name = serializer.validated_data['name']
        unit.save()
        
        logger.info(f"User {request.user.username} updated unit {unit.name}")
        
        return success_response(data=UnitSerializer(unit).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除单位"""
        try:
            unit = Unit.objects.get(pk=pk)
        except Unit.DoesNotExist:
            return error_response(message='单位不存在', code=404)
        
        if unit.is_linked:
            return error_response(message='该单位已被关联，无法删除')
        
        name = unit.name
        unit.delete()
        
        logger.info(f"User {request.user.username} deleted unit {name}")
        
        return success_response(message='删除成功')


class UnitBatchDeleteView(APIView):
    """单位批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的单位')
        
        # 只删除未关联的单位
        units = Unit.objects.filter(pk__in=ids)
        deleted_count = 0
        for unit in units:
            if not unit.is_linked:
                unit.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} units")
        
        return success_response(message=f'成功删除 {deleted_count} 个单位')


class UnitAllView(APIView):
    """获取所有单位（用于下拉选择）"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        units = Unit.objects.filter(is_active=True).order_by('name')
        serializer = UnitSerializer(units, many=True)
        return success_response(data=serializer.data)


# ==================== 品类管理 ====================

class CategoryListView(APIView):
    """品类列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Category.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        categories = queryset[start:end]
        
        serializer = CategorySerializer(categories, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建品类"""
        serializer = CategoryCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit = Unit.objects.get(pk=serializer.validated_data['unit'])
        category = Category.objects.create(
            name=serializer.validated_data['name'],
            unit=unit,
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created category {category.name}")
        
        return success_response(data=CategorySerializer(category).data, message='创建成功')


class CategoryDetailView(APIView):
    """品类详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新品类"""
        try:
            category = Category.objects.get(pk=pk)
        except Category.DoesNotExist:
            return error_response(message='品类不存在', code=404)
        
        serializer = CategoryCreateSerializer(data=request.data, context={'instance': category})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        category.name = serializer.validated_data['name']
        category.unit = Unit.objects.get(pk=serializer.validated_data['unit'])
        category.save()
        
        logger.info(f"User {request.user.username} updated category {category.name}")
        
        return success_response(data=CategorySerializer(category).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除品类"""
        try:
            category = Category.objects.get(pk=pk)
        except Category.DoesNotExist:
            return error_response(message='品类不存在', code=404)
        
        if category.is_linked:
            return error_response(message='该品类已被关联，无法删除')
        
        name = category.name
        category.delete()
        
        logger.info(f"User {request.user.username} deleted category {name}")
        
        return success_response(message='删除成功')


class CategoryBatchDeleteView(APIView):
    """品类批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的品类')
        
        categories = Category.objects.filter(pk__in=ids)
        deleted_count = 0
        for category in categories:
            if not category.is_linked:
                category.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} categories")
        
        return success_response(message=f'成功删除 {deleted_count} 个品类')


class CategoryAllView(APIView):
    """获取所有品类（用于下拉选择）"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        categories = Category.objects.filter(is_active=True).order_by('name')
        serializer = CategorySerializer(categories, many=True)
        return success_response(data=serializer.data)


# ==================== 品种管理 ====================

class VarietyListView(APIView):
    """品种列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Variety.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        varieties = queryset[start:end]
        
        serializer = VarietySerializer(varieties, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建品种"""
        serializer = VarietyCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))
        
        category = Category.objects.get(pk=serializer.validated_data['category'])
        variety = Variety.objects.create(
            name=serializer.validated_data['name'],
            category=category,
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created variety {variety.name}")
        
        return success_response(data=VarietySerializer(variety).data, message='创建成功')


class VarietyDetailView(APIView):
    """品种详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新品种"""
        try:
            variety = Variety.objects.get(pk=pk)
        except Variety.DoesNotExist:
            return error_response(message='品种不存在', code=404)
        
        serializer = VarietyCreateSerializer(data=request.data, context={'instance': variety})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))
        
        variety.name = serializer.validated_data['name']
        variety.category = Category.objects.get(pk=serializer.validated_data['category'])
        variety.save()
        
        logger.info(f"User {request.user.username} updated variety {variety.name}")
        
        return success_response(data=VarietySerializer(variety).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除品种"""
        try:
            variety = Variety.objects.get(pk=pk)
        except Variety.DoesNotExist:
            return error_response(message='品种不存在', code=404)
        
        if variety.is_in_stock:
            return error_response(message='该品种已入库，无法删除')
        
        name = variety.name
        variety.delete()
        
        logger.info(f"User {request.user.username} deleted variety {name}")
        
        return success_response(message='删除成功')


class VarietyBatchDeleteView(APIView):
    """品种批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的品种')
        
        varieties = Variety.objects.filter(pk__in=ids)
        deleted_count = 0
        for variety in varieties:
            if not variety.is_in_stock:
                variety.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} varieties")
        
        return success_response(message=f'成功删除 {deleted_count} 个品种')


class VarietyTemplateView(APIView):
    """品种导入模板下载"""
    permission_classes = []  # 允许匿名访问，通过token参数验证
    
    def get(self, request):
        # 从URL参数获取token进行验证
        from apps.authentication.backends import decode_token
        from apps.authentication.models import User
        
        token = request.query_params.get('token')
        if not token:
            return error_response(message='缺少认证信息', code=401)
        
        payload = decode_token(token)
        if not payload:
            return error_response(message='认证信息无效或已过期', code=401)
        
        try:
            user = User.objects.get(pk=payload['user_id'])
        except User.DoesNotExist:
            return error_response(message='用户不存在', code=401)
        
        wb = Workbook()
        
        # 第一个表格 - 导入模板
        ws1 = wb.active
        ws1.title = '品种导入'
        
        # 设置表头样式
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        headers = ['品种', '品类', '单位']
        for col, header in enumerate(headers, 1):
            cell = ws1.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 设置列宽
        ws1.column_dimensions['A'].width = 25
        ws1.column_dimensions['B'].width = 20
        ws1.column_dimensions['C'].width = 15
        
        # 第二个表格 - 品类参考
        ws2 = wb.create_sheet(title='品类参考')
        
        headers2 = ['品类', '单位']
        for col, header in enumerate(headers2, 1):
            cell = ws2.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 填充品类数据
        categories = Category.objects.filter(is_active=True).select_related('unit')
        for row, category in enumerate(categories, 2):
            ws2.cell(row=row, column=1, value=category.name).border = thin_border
            ws2.cell(row=row, column=2, value=category.unit.name).border = thin_border
        
        ws2.column_dimensions['A'].width = 20
        ws2.column_dimensions['B'].width = 15
        
        # 返回Excel文件
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=variety_import_template.xlsx'
        
        return response


class VarietyImportView(APIView):
    """品种导入视图"""
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]
    
    def post(self, request):
        if 'file' not in request.FILES:
            return error_response(message='请上传文件')
        
        file = request.FILES['file']
        
        try:
            wb = load_workbook(file)
            ws = wb.active
        except Exception as e:
            return error_response(message='文件格式错误，请上传Excel文件')
        
        # 获取所有品类及其单位
        categories = {c.name: c for c in Category.objects.filter(is_active=True).select_related('unit')}
        
        can_import = []
        cannot_import = []
        
        for row in range(2, ws.max_row + 1):
            variety_name = ws.cell(row=row, column=1).value
            category_name = ws.cell(row=row, column=2).value
            unit_name = ws.cell(row=row, column=3).value
            
            if not variety_name:
                continue
            
            variety_name = str(variety_name).strip()
            category_name = str(category_name).strip() if category_name else ''
            unit_name = str(unit_name).strip() if unit_name else ''
            
            # 验证
            error_msg = None
            
            if not variety_name:
                error_msg = '品种名称不能为空'
            elif len(variety_name) > 20:
                error_msg = '品种名称最多20个字'
            elif not category_name:
                error_msg = '品类不能为空'
            elif category_name not in categories:
                error_msg = f'品类"{category_name}"不存在'
            elif not unit_name:
                error_msg = '单位不能为空'
            elif categories.get(category_name) and categories[category_name].unit.name != unit_name:
                error_msg = f'单位与品类不匹配，应为"{categories[category_name].unit.name}"'
            elif Variety.objects.filter(name=variety_name, category__name=category_name).exists():
                error_msg = '该品种已存在'
            
            if error_msg:
                cannot_import.append({
                    'row': row,
                    'variety': variety_name,
                    'category': category_name,
                    'unit': unit_name,
                    'reason': error_msg
                })
            else:
                can_import.append({
                    'row': row,
                    'variety': variety_name,
                    'category': category_name,
                    'unit': unit_name
                })
        
        # 如果是预览请求
        if request.data.get('preview') == 'true':
            return success_response(data={
                'can_import': can_import,
                'cannot_import': cannot_import,
                'can_import_count': len(can_import),
                'cannot_import_count': len(cannot_import)
            })
        
        # 执行导入
        imported_count = 0
        for item in can_import:
            category = categories[item['category']]
            Variety.objects.create(
                name=item['variety'],
                category=category,
                created_by=request.user
            )
            imported_count += 1
        
        logger.info(f"User {request.user.username} imported {imported_count} varieties")
        
        return success_response(
            data={
                'imported_count': imported_count,
                'failed_count': len(cannot_import),
                'failed_items': cannot_import
            },
            message=f'成功导入 {imported_count} 个品种'
        )


# ==================== 其他视图占位 ====================

class DashboardView(APIView):
    """仪表盘视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'message': '仪表盘功能开发中...'
        })


class GoodsListView(APIView):
    """货物列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class StockInListView(APIView):
    """入库记录列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class StockOutListView(APIView):
    """出库记录列表/申请视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = StockOut.objects.select_related('goods').order_by('-created_at')

        status = request.query_params.get('status')
        if status:
            queryset = queryset.filter(status=status)

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size

        total = queryset.count()
        items = list(queryset[start:end])
        data = StockOutSerializer(items, many=True).data
        # 附带逾期放行提示，便于前端在审批前识别
        for item, obj in zip(data, items):
            item['requires_extra_approval'] = review_service.goods_review_overdue(
                obj.goods, on_date=timezone.localdate(obj.created_at)
            )
            item['extra_approval_id'] = obj.extra_approval_id

        return success_response(data={
            'list': data,
            'total': total,
            'page': page,
            'page_size': page_size
        })

    def post(self, request):
        """申请出库；物资复核逾期时返回需额外审批标记，不阻断申请。"""
        goods_id = request.data.get('goods')
        quantity = request.data.get('quantity')
        receiver = (request.data.get('receiver') or '').strip()

        if not goods_id or not quantity or not receiver:
            return error_response(message='物资、数量和领用人不能为空')
        try:
            goods = Goods.objects.get(pk=goods_id)
        except Goods.DoesNotExist:
            return error_response(message='物资不存在', code=404)
        try:
            quantity = Decimal(str(quantity))
        except Exception:
            return error_response(message='出库数量格式不正确')
        if quantity <= 0:
            return error_response(message='出库数量必须大于0')
        if quantity > goods.quantity:
            return error_response(message='出库数量不能超过当前库存')

        stock_out = StockOut.objects.create(
            goods=goods,
            operator=request.user,
            receiver=receiver,
            receiver_dept=request.data.get('receiver_dept', ''),
            quantity=quantity,
            status='pending',
        )

        data = StockOutSerializer(stock_out).data
        data['requires_extra_approval'] = stock_out.requires_extra_approval
        return success_response(data=data, message='出库申请已提交')


class StockOutReleaseView(APIView):
    """出库放行：逾期物资必须先通过额外审批"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        try:
            stock_out = StockOut.objects.select_related('goods').get(pk=pk)
        except StockOut.DoesNotExist:
            return error_response(message='出库记录不存在', code=404)

        if stock_out.status != 'approved':
            return error_response(message='出库单未经审批通过，不能放行')

        try:
            review_service.assert_release_allowed(stock_out)
        except PermissionError as exc:
            return error_response(message=str(exc))

        stock_out.status = 'completed'
        stock_out.stock_out_time = timezone.now()
        stock_out.save(update_fields=['status', 'stock_out_time'])

        # 扣减库存
        goods = stock_out.goods
        goods.quantity -= stock_out.quantity
        goods.save(update_fields=['quantity', 'updated_at'])

        logger.info("出库放行: %s x%s（操作人 %s）",
                    goods.name, stock_out.quantity, request.user.username)
        return success_response(data=StockOutSerializer(stock_out).data, message='放行成功')


class StockOutApproveView(APIView):
    """出库常规审批"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        try:
            stock_out = StockOut.objects.get(pk=pk)
        except StockOut.DoesNotExist:
            return error_response(message='出库记录不存在', code=404)

        if not request.user.is_admin:
            return error_response(message='仅管理员可审批出库申请', code=403)
        if stock_out.status != 'pending':
            return error_response(message='该出库单已审批')

        action = request.data.get('action', 'approve')
        approval_status = 'approved' if action == 'approve' else 'rejected'
        approval = Approval.objects.create(
            stock_out=stock_out,
            approver=request.user,
            approval_type='stock_out',
            status=approval_status,
            remark=request.data.get('remark', ''),
        )
        stock_out.status = approval_status
        stock_out.save(update_fields=['status'])

        return success_response(data=ApprovalSerializer(approval).data,
                                message='审批通过' if action == 'approve' else '已拒绝')


class StockOutExtraApprovalView(APIView):
    """逾期物资放行前的额外审批"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        try:
            stock_out = StockOut.objects.select_related('goods').get(pk=pk)
        except StockOut.DoesNotExist:
            return error_response(message='出库记录不存在', code=404)

        if not stock_out.requires_extra_approval:
            return error_response(message='该物资不存在复核逾期，无需额外审批')
        if stock_out.extra_approval_id is not None:
            return error_response(message='该出库单已完成逾期额外审批')
        if not request.user.is_admin:
            return error_response(message='仅管理员可进行逾期放行额外审批', code=403)

        approval = review_service.approve_overdue_release(
            stock_out=stock_out,
            approver=request.user,
            remark=request.data.get('remark', ''),
        )
        return success_response(data=ApprovalSerializer(approval).data, message='额外审批通过')


# ==================== 封存复核 ====================

class ReviewRuleListView(APIView):
    """复核规则版本列表 / 发布新版本"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = ReviewRuleVersion.objects.select_related('category').order_by(
            '-created_at'
        )
        category_id = request.query_params.get('category')
        status = request.query_params.get('status')
        if category_id:
            queryset = queryset.filter(category_id=category_id)
        if status:
            queryset = queryset.filter(status=status)

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size

        return success_response(data={
            'list': ReviewRuleSerializer(queryset[start:end], many=True).data,
            'total': queryset.count(),
            'page': page,
            'page_size': page_size
        })

    def post(self, request):
        if not request.user.is_admin:
            return error_response(message='仅管理员可发布复核规则', code=403)

        serializer = ReviewRulePublishSerializer(data=request.data)
        if not serializer.is_valid():
            first_error = list(serializer.errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))

        data = serializer.validated_data
        category = None
        if data.get('category'):
            category = Category.objects.get(pk=data['category'])

        rule = review_service.publish_rule(
            category=category,
            name=data['name'],
            interval_days=data['interval_days'],
            check_items=data['check_items'],
            overdue_grace_days=data.get('overdue_grace_days', 0),
            published_by=request.user,
        )
        return success_response(data=ReviewRuleSerializer(rule).data, message='规则发布成功')


class ReviewTaskListView(APIView):
    """复核待办列表"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = ReviewTask.objects.select_related(
            'goods', 'goods__variety__category', 'rule_version'
        ).order_by('due_date', '-id')

        status = request.query_params.get('status', 'pending')
        overdue = request.query_params.get('overdue')
        if status:
            queryset = queryset.filter(status=status)
        if overdue == 'true':
            today = timezone.localdate()
            task_ids = [
                t.id for t in queryset.select_related('rule_version')
                if today > t.due_date + datetime.timedelta(days=t.rule_version.overdue_grace_days)
            ]
            queryset = queryset.filter(id__in=task_ids)

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size

        return success_response(data={
            'list': ReviewTaskSerializer(queryset[start:end], many=True).data,
            'total': queryset.count(),
            'page': page,
            'page_size': page_size
        })


class ReviewTaskGenerateView(APIView):
    """批量生成复核待办（重复执行幂等）"""
    permission_classes = [IsAuthenticated]

    def post(self, request):
        result = review_service.generate_review_tasks()
        return success_response(data={
            'created_count': result['created_count'],
            'skipped_count': result['skipped_count'],
        }, message=f"新增 {result['created_count']} 条待办")


class ReviewTaskCompleteView(APIView):
    """登记复核结果"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        try:
            task = ReviewTask.objects.select_related('rule_version').get(pk=pk)
        except ReviewTask.DoesNotExist:
            return error_response(message='复核待办不存在', code=404)

        serializer = ReviewCompleteSerializer(data=request.data)
        if not serializer.is_valid():
            first_error = list(serializer.errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))

        data = serializer.validated_data
        try:
            record = review_service.complete_review_task(
                task=task,
                conclusion=data['conclusion'],
                reviewer=request.user,
                abnormal_items=data.get('abnormal_items') or [],
                review_date=data.get('review_date'),
                remark=data.get('remark', ''),
            )
        except ValueError as exc:
            return error_response(message=str(exc))

        return success_response(data=ReviewRecordSerializer(record).data, message='复核结果已登记')


class ReviewRecordListView(APIView):
    """复核结果列表（历史结论）"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = ReviewRecord.objects.select_related(
            'goods', 'rule_version', 'reviewer'
        ).order_by('-review_date', '-id')

        goods_id = request.query_params.get('goods')
        if goods_id:
            queryset = queryset.filter(goods_id=goods_id)

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size

        return success_response(data={
            'list': ReviewRecordSerializer(queryset[start:end], many=True).data,
            'total': queryset.count(),
            'page': page,
            'page_size': page_size
        })


class WarningListView(APIView):
    """预警记录列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = Warning.objects.select_related('goods').order_by('-created_at')

        warning_type = request.query_params.get('type')
        if warning_type:
            queryset = queryset.filter(type=warning_type)
        is_read = request.query_params.get('is_read')
        if is_read is not None:
            queryset = queryset.filter(is_read=is_read in ('true', '1', 'True'))

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size

        return success_response(data={
            'list': WarningSerializer(queryset[start:end], many=True).data,
            'total': queryset.count(),
            'page': page,
            'page_size': page_size
        })


class ApprovalListView(APIView):
    """审批记录列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = Approval.objects.select_related(
            'stock_out', 'approver'
        ).order_by('-created_at')

        approval_type = request.query_params.get('approval_type')
        if approval_type:
            queryset = queryset.filter(approval_type=approval_type)

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size

        return success_response(data={
            'list': ApprovalSerializer(queryset[start:end], many=True).data,
            'total': queryset.count(),
            'page': page,
            'page_size': page_size
        })
